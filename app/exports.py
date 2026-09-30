"""Deliverables. Every report follows one standard structure; its content stays dynamic.

    1 Introduction · 2 Benchmarking Methodology · 3 Benchmark Models ·
    4 Comparative Analysis · 5 Recommendations & Conclusions · 6 References

Only approved content is used. Citations are APA 7 (in-text + reference list) built from accepted evidence.
"""
import io
import re

import method as bm
import scoring
from db import q, one, jl

TEAL, DEEP, MINT, MINT2, BERRY, GREY = "056073", "0B2F38", "D1F1ED", "E7F8F6", "A11249", "5B6770"

REPORT_OUTLINE = [
    ("Introduction", []),
    ("Benchmarking Methodology", []),
    ("Benchmark Models", ["profile", "assessment"]),
    ("Comparative Analysis", ["comparison"]),
    ("Recommendations & Conclusions", ["lesson", "recommendation", "limitation"]),
    ("References", []),
]
REPORT_SECTIONS_REQUIRED = [("Benchmark Models", ["profile", "assessment"]), ("Comparative Analysis", ["comparison"]),
                            ("Recommendations & Conclusions", ["recommendation"])]
TYPE_LABEL = {"verified_fact": "", "benchmark_comparison": "Benchmark finding", "ai_synthesis": "Synthesis",
              "ai_interpretation": "Interpretation", "client_implication": "Recommendation", "research_gap": "Limitation"}


def _clean(t):
    return re.sub(r"\s*\[(Interpretation|Synthesis)\]\s*", " ", t or "").strip()


def references(eid, used_only=True):
    return [{"ref": r["ref"], "intext": r["intext"], "codes": r["codes"], "url": r["url"]} for r in bm.reference_list(eid, used_only)]


def data(eid):
    """Everything the report needs, gathered once."""
    e = one("SELECT e.*, u.name lead FROM engagements e LEFT JOIN users u ON u.id=e.lead_id WHERE e.id=?", (eid,))
    items = q("SELECT * FROM content_items WHERE engagement_id=? AND status IN ('approved','open_gap') AND in_deliverable=1 ORDER BY id", (eid,))
    for i in items:
        i["evidence_ids"] = jl(i["evidence_ids"])
    m = scoring.matrix(eid)
    cmap = bm.citations(eid)
    stats = one("SELECT (SELECT COUNT(*) FROM research_tasks WHERE engagement_id=?) tasks,"
                " (SELECT COUNT(*) FROM research_tasks WHERE engagement_id=? AND status='gap') gaps,"
                " (SELECT COUNT(*) FROM evidence WHERE engagement_id=? AND status='accepted') accepted,"
                " (SELECT COUNT(*) FROM evidence WHERE engagement_id=? AND status='rejected') rejected,"
                " (SELECT COUNT(*) FROM search_log s JOIN research_tasks t ON t.id=s.task_id WHERE t.engagement_id=?) searches",
                (eid,) * 5)
    return {"e": e, "scope": jl(e["scope"], {}), "questions": [x.strip(" -•") for x in (e["key_questions"] or "").splitlines() if x.strip(" -•")],
            "story": jl(e["storyline"], {}), "items": items, "m": m, "cmap": cmap, "stats": stats,
            "released": one("SELECT status FROM gates WHERE engagement_id=? AND stage=6", (eid,))["status"] == "approved",
            "rules": q("SELECT * FROM source_rules ORDER BY sort"), "refs": references(eid)}


def outline(eid):
    d = data(eid)
    out = []
    for name, secs in REPORT_OUTLINE:
        its = [i for i in d["items"] if i["section"] in secs]
        auto = {"Introduction": "Brief, scope and key questions", "Benchmarking Methodology": "Analysis method, framework, benchmark set and source policy",
                "References": f"{len(d['refs'])} APA references generated from cited evidence"}.get(name)
        out.append({"section": name, "items": len(its), "auto": auto,
                    "ready": bool(auto) or bool(its) or name not in [x[0] for x in REPORT_SECTIONS_REQUIRED]})
    return {"sections": out, "story": d["story"], "released": d["released"]}


def _cite(item, cmap):
    return bm.intext(item["evidence_ids"], cmap)


def _para_text(item, cmap):
    label = TYPE_LABEL.get(item["content_type"], "")
    txt = _clean(item["text"])
    c = _cite(item, cmap)
    return label, (txt[:-1] if txt.endswith(".") and c else txt) + (f" {c}." if c else "")


# ------------------------------------------------------------------ Word report
def docx(eid):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm

    d = data(eid)
    e, m, cmap = d["e"], d["m"], d["cmap"]
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(10.5)
    for lvl, size in ((1, 16), (2, 12.5), (3, 11)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name, h.font.size, h.font.color.rgb, h.font.bold = "Arial", Pt(size), RGBColor.from_string(TEAL if lvl < 3 else DEEP), True
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor)
        tcPr.append(sh)

    def table(head, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(head):
            c = t.rows[0].cells[i]; c.text = ""
            r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(255, 255, 255)
            shade(c, TEAL)
        for k, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(v if v is not None else "—")); run.font.size = Pt(9)
                if k % 2:
                    shade(cells[i], MINT2)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        doc.add_paragraph()
        return t

    def item_par(it):
        label, txt = _para_text(it, cmap)
        p = doc.add_paragraph(style="List Bullet")
        if it.get("title") and it["section"] not in ("comparison",):
            p.add_run(it["title"] + ". ").bold = True
        if label:
            r = p.add_run(f"[{label}] "); r.italic = True; r.font.color.rgb = RGBColor.from_string(GREY)
        p.add_run(txt)
        if it.get("assumptions"):
            r = p.add_run(f" Conditions: {it['assumptions']}"); r.italic = True

    # cover
    p = doc.add_paragraph(); r = p.add_run("MUSHAR CONSULTING · BENCHMARKING REPORT"); r.font.size = Pt(9); r.bold = True; r.font.color.rgb = RGBColor.from_string(TEAL)
    p = doc.add_paragraph(); r = p.add_run(e["title"]); r.font.size = Pt(24); r.bold = True; r.font.color.rgb = RGBColor.from_string(DEEP)
    doc.add_paragraph(f"{e['client']}  ·  {e['code']}  ·  Engagement lead: {e['lead']}")
    if not d["released"]:
        r = doc.add_paragraph().add_run("DRAFT — not yet released through QA"); r.bold = True; r.font.color.rgb = RGBColor.from_string(BERRY)
    if d["story"].get("executive_summary"):
        doc.add_heading("Executive summary", 2)
        for x in d["story"]["executive_summary"]:
            c = bm.intext(x.get("evidence_ids") or [], cmap)
            doc.add_paragraph(_clean(x["text"]) + (f" {c}" if c else ""), style="List Bullet")
    doc.add_page_break()

    # 1 Introduction
    doc.add_heading("1. Introduction", 1)
    doc.add_heading("Purpose and objectives", 2); doc.add_paragraph(e["objective"] or "—")
    if e["context"]:
        doc.add_heading("Engagement context", 2); doc.add_paragraph(e["context"])
    doc.add_heading("Decision this benchmark supports", 2); doc.add_paragraph(e["decision_statement"] or "—")
    if e["requirements"]:
        doc.add_heading("Client requirements", 2); doc.add_paragraph(e["requirements"])
    if d["questions"]:
        doc.add_heading("Key questions", 2)
        for x in d["questions"]:
            doc.add_paragraph(x, style="List Number")
    sc = d["scope"]
    doc.add_heading("Scope", 2)
    table(["Scope element", "Definition"], [(k.capitalize(), sc.get(k)) for k in ("inclusions", "exclusions", "geography", "period", "constraints") if sc.get(k)], [4, 12])
    if e["expected_outcomes"]:
        doc.add_heading("Expected outcomes", 2); doc.add_paragraph(e["expected_outcomes"])

    # 2 Methodology
    doc.add_heading("2. Benchmarking Methodology", 1)
    doc.add_heading("Analysis approach", 2)
    doc.add_paragraph(bm.ANALYSIS.get(e["analysis_method"] or "mixed", ""))
    if e["analysis_rationale"]:
        doc.add_paragraph(e["analysis_rationale"])
    doc.add_heading("Assessment framework", 2)
    rows = []
    for dm in m["dimensions"]:
        for c in [c for c in m["criteria"] if c["dimension_id"] == dm["id"]]:
            rows.append((dm["name"], c["name"], bm.ASSESSMENT.get(c["assessment_type"], (c["assessment_type"],))[0],
                         "Yes" if c["is_scored"] else "No", c["question"]))
    table(["Dimension", "Criterion", "Assessment method", "Scored", "Research question"], rows, [3, 3.2, 3.2, 1.3, 5.8])
    doc.add_heading("Benchmark set", 2)
    table(["Benchmark", "Type", "Role", "Selection rationale"],
          [(p["name"], bm.KINDS.get(p["kind"], p["kind"]), bm.ROLES.get(p["role"], (p["role"] or "—",))[0], p["rationale"]) for p in m["comparators"]], [4, 2.3, 3.3, 7])
    doc.add_heading("Research and evidence standards", 2)
    doc.add_paragraph("Research used credible, publicly accessible sources, prioritising primary and original publications. "
                      "Every material finding is traceable to accepted evidence, cited in APA style.")
    for r in d["rules"]:
        if r["priority"] != "REJECT":
            doc.add_paragraph(f"{r['priority']} — {r['category']}: {r['examples']}", style="List Bullet")
    stt = d["stats"]
    doc.add_paragraph(f"Research coverage: {stt['tasks']} benchmark × question tasks, {stt['searches']} logged searches, "
                      f"{stt['accepted']} accepted evidence items ({stt['rejected']} sources rejected on review), {stt['gaps']} evidence gap(s) disclosed.")
    if m["scoring"]:
        doc.add_paragraph("Scoring: scored criteria are normalised to 0–100 and weighted. Unknown values are excluded rather than scored as zero; "
                          f"a composite score is shown only when coverage is at least {m['threshold']}%. Quantitative values are scored only after "
                          "definition, period, unit and denominator checks pass.")
    else:
        doc.add_paragraph("No numerical scoring was applied: the questions are best answered through descriptive comparison and practice analysis.")

    # 3 Benchmark models
    doc.add_heading("3. Benchmark Models", 1)
    for p in m["comparators"]:
        its = [i for i in d["items"] if i["comparator_id"] == p["id"] and i["section"] in ("profile", "assessment")]
        doc.add_heading(p["name"], 2)
        r = doc.add_paragraph().add_run(f"{bm.ROLES.get(p['role'], ('',))[0]} · {bm.KINDS.get(p['kind'], p['kind'] or '')} · {p['region'] or ''}")
        r.italic = True; r.font.color.rgb = RGBColor.from_string(GREY)
        for it in its:
            item_par(it)
        if not its:
            doc.add_paragraph("No approved content for this benchmark.")

    # 4 Comparative analysis
    doc.add_heading("4. Comparative Analysis", 1)
    head = ["Criterion"] + [p["name"].split(" — ")[0] for p in m["comparators"]]
    rows = [[c["name"]] + [m["cells"][f"{c['id']}:{p['id']}"]["display"] for p in m["comparators"]] for c in m["criteria"]]
    if m["scoring"]:
        rows.append(["Overall score (coverage)"] + [f"{m['scores'][p['id']]['overall'] if m['scores'][p['id']]['overall'] is not None else 'n/a'} ({m['scores'][p['id']]['coverage']}%)" for p in m["comparators"]])
    table(head, rows)
    doc.add_heading("Findings", 2)
    for it in [i for i in d["items"] if i["section"] == "comparison"]:
        item_par(it)

    # 5 Recommendations & conclusions
    doc.add_heading("5. Recommendations & Conclusions", 1)
    for sec, title in (("lesson", "Lessons learned"), ("recommendation", "Recommendations"), ("limitation", "Limitations")):
        its = [i for i in d["items"] if i["section"] == sec]
        if its:
            doc.add_heading(title, 2)
            for it in its:
                item_par(it)
    if d["story"].get("conclusions"):
        doc.add_heading("Conclusions", 2)
        for x in d["story"]["conclusions"]:
            c = bm.intext(x.get("evidence_ids") or [], cmap)
            doc.add_paragraph(_clean(x["text"]) + (f" {c}" if c else ""), style="List Bullet")

    # 6 References (APA 7, hanging indent)
    doc.add_heading("6. References", 1)
    for r in d["refs"]:
        p = doc.add_paragraph(r["ref"])
        p.paragraph_format.left_indent = Cm(1.27); p.paragraph_format.first_line_indent = Cm(-1.27)
    if not d["refs"]:
        doc.add_paragraph("No sources cited yet.")
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PowerPoint (same structure)
def pptx(eid):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    d = data(eid)
    e, m, cmap = d["e"], d["m"], d["cmap"]
    rgb = lambda h: RGBColor.from_string(h)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    n = [0]

    def box(s, x, y, w, h, fill=None):
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill:
            sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        sh.line.fill.background(); sh.shadow.inherit = False
        return sh

    def text(s, x, y, w, h, t, size=12, color=DEEP, bold=False):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        for i, line in enumerate(t if isinstance(t, list) else [t]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run(); r.text = line
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Arial"
            r.font.color.rgb = rgb(color); p.space_after = Pt(6)
        return tb

    def slide(title, kicker=""):
        s = prs.slides.add_slide(blank); n[0] += 1
        box(s, 0, 0, 13.333, 0.12, TEAL)
        if kicker:
            text(s, 0.6, 0.35, 12, 0.3, kicker.upper(), 10, TEAL, True)
        text(s, 0.6, 0.6, 12.1, 0.9, title, 24, DEEP, True)
        text(s, 0.6, 7.0, 9, 0.3, f"{e['code']} · {e['title']}" + ("" if d["released"] else " · DRAFT — not released"), 8, GREY)
        text(s, 12.2, 7.0, 0.6, 0.3, str(n[0]), 8, GREY)
        return s

    def divider(num, name):
        s = prs.slides.add_slide(blank); n[0] += 1
        box(s, 0, 0, 13.333, 7.5, DEEP)
        text(s, 0.95, 2.6, 3, 1, f"{num:02d}", 54, "8EDDD2", True)
        text(s, 0.95, 3.7, 11, 1, name, 32, "FFFFFF", True)

    def lines_for(its, size_hint=None):
        out = []
        for it in its:
            label, txt = _para_text(it, cmap)
            out.append(f"■ {(it['title'] + ': ') if it['title'] and it['section'] != 'comparison' else ''}{('[' + label + '] ') if label else ''}{txt}")
        return out or ["No approved content."]

    def paged(title, kicker, lines, per=6, size=13):
        for k in range(0, max(len(lines), 1), per):
            s = slide(title + (" (cont.)" if k else ""), kicker)
            text(s, 0.6, 1.7, 12.1, 5.2, lines[k:k + per] or ["—"], size)

    # title
    s = prs.slides.add_slide(blank); n[0] += 1
    box(s, 0, 0, 13.333, 7.5, DEEP); box(s, 0.6, 2.2, 0.12, 2.2, "8EDDD2")
    text(s, 0.95, 2.1, 11, 1.4, e["title"], 36, "FFFFFF", True)
    text(s, 0.95, 3.5, 11, 0.6, e["client"], 18, "8EDDD2")
    text(s, 0.95, 4.1, 11, 0.5, "Benchmarking report · Mushar Consulting" + ("" if d["released"] else " · DRAFT"), 12, "D1F1ED")
    if d["story"].get("executive_summary"):
        paged("Executive summary", "Summary", [f"■ {_clean(x['text'])} {bm.intext(x.get('evidence_ids') or [], cmap)}".strip() for x in d["story"]["executive_summary"]], 6, 14)

    # 1 Introduction
    divider(1, "Introduction")
    s = slide("Purpose, decision and key questions", "Introduction")
    text(s, 0.6, 1.6, 5.9, 5.2, ["Objective", e["objective"] or "—", "", "Decision supported", e["decision_statement"] or "—"], 12)
    text(s, 6.9, 1.6, 5.8, 5.2, ["Key questions"] + [f"{i + 1}. {x}" for i, x in enumerate(d["questions"])], 12)

    # 2 Methodology
    divider(2, "Benchmarking Methodology")
    s = slide("Approach, framework and benchmark set", "Methodology")
    text(s, 0.6, 1.5, 12.1, 0.7, bm.ANALYSIS.get(e["analysis_method"] or "mixed", ""), 12, TEAL, True)
    y = 2.2
    for dm in m["dimensions"]:
        cr = [c for c in m["criteria"] if c["dimension_id"] == dm["id"]]
        box(s, 0.6, y, 3.0, 0.3 + 0.26 * len(cr), MINT)
        text(s, 0.7, y + 0.02, 2.8, 0.3, dm["name"], 11, DEEP, True)
        text(s, 3.8, y, 3.9, 0.26 * len(cr) + 0.3, [f"{c['name']} — {bm.ASSESSMENT.get(c['assessment_type'], (c['assessment_type'],))[0]}" for c in cr], 9)
        y += 0.45 + 0.26 * len(cr)
    text(s, 8.0, 2.2, 4.8, 4.6, ["Benchmark set"] + [f"• {p['name'].split(' — ')[0]} — {bm.ROLES.get(p['role'], ('—',))[0]}" for p in m["comparators"]], 10)

    # 3 Benchmark models
    divider(3, "Benchmark Models")
    for p in m["comparators"]:
        its = [i for i in d["items"] if i["comparator_id"] == p["id"] and i["section"] in ("profile", "assessment")]
        paged(p["name"], f"Benchmark model · {bm.ROLES.get(p['role'], ('',))[0]}", lines_for(its), 5)

    # 4 Comparative analysis
    divider(4, "Comparative Analysis")
    s = slide("Comparison across the framework", "Comparative analysis")
    comps = m["comparators"]
    rows, cols = len(m["criteria"]) + 1 + (1 if m["scoring"] else 0), len(comps) + 1
    tbl = s.shapes.add_table(rows, cols, Inches(0.6), Inches(1.6), Inches(12.1), Inches(0.34 * rows)).table
    tbl.columns[0].width = Inches(3.4)
    for j in range(1, cols):
        tbl.columns[j].width = Inches(8.7 / max(cols - 1, 1))

    def cell(r, c, t, fill=None, bold=False, color=DEEP):
        ce = tbl.cell(r, c); ce.text = t
        for run in ce.text_frame.paragraphs[0].runs:
            run.font.size, run.font.bold, run.font.name = Pt(9), bold, "Arial"
            run.font.color.rgb = rgb(color)
        if fill:
            ce.fill.solid(); ce.fill.fore_color.rgb = rgb(fill)
    cell(0, 0, "Criterion", TEAL, True, "FFFFFF")
    for j, p in enumerate(comps, 1):
        cell(0, j, p["name"].split(" — ")[0], TEAL, True, "FFFFFF")
    shade = lambda x: ("8EDDD2" if x["norm"] >= .75 else "D1F1ED" if x["norm"] >= .5 else "F7F8EB" if x["norm"] >= .25 else "F3D6E0") if x["norm"] is not None \
        else {"pos": "8EDDD2", "mid": "D1F1ED", "neutral": "F7F8EB"}.get(x["tone"], "FFFFFF")
    for i, c in enumerate(m["criteria"], 1):
        cell(i, 0, c["name"])
        for j, p in enumerate(comps, 1):
            x = m["cells"][f"{c['id']}:{p['id']}"]
            cell(i, j, x["display"], shade(x))
    if m["scoring"]:
        cell(rows - 1, 0, "Overall score (coverage)", MINT2, True)
        for j, p in enumerate(comps, 1):
            sc = m["scores"][p["id"]]
            cell(rows - 1, j, f"{sc['overall'] if sc['overall'] is not None else 'n/a'} ({sc['coverage']}%)", MINT2, True)
    paged("Comparative findings", "Comparative analysis", lines_for([i for i in d["items"] if i["section"] == "comparison"]), 4, 12)

    # 5 Recommendations & conclusions
    divider(5, "Recommendations & Conclusions")
    for sec, title in (("lesson", "Lessons learned"), ("recommendation", "Recommendations"), ("limitation", "Limitations")):
        its = [i for i in d["items"] if i["section"] == sec]
        if its:
            paged(title, "Recommendations & conclusions", lines_for(its), 5)
    if d["story"].get("conclusions"):
        paged("Conclusions", "Recommendations & conclusions",
              [f"■ {_clean(x['text'])} {bm.intext(x.get('evidence_ids') or [], cmap)}".strip() for x in d["story"]["conclusions"]], 5)

    # 6 References
    divider(6, "References")
    refs = [r["ref"] for r in d["refs"]]
    paged("References", "APA 7", refs or ["No sources cited yet."], 12, 9)
    buf = io.BytesIO(); prs.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ Excel
def xlsx(eid):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    head = PatternFill("solid", fgColor=TEAL)

    def sheet(ws, cols, rows):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for r in rows:
            ws.append([r.get(k) if isinstance(r, dict) else r[i] for i, k in enumerate(cols)])
        for i, col in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(60, max(12, len(col) + 4))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    cmap = bm.citations(eid)
    ev = q("SELECT e.*, p.name AS comparator, c.name AS criterion FROM evidence e JOIN research_tasks t ON t.id=e.task_id "
           "JOIN comparators p ON p.id=t.comparator_id JOIN questions qu ON qu.id=t.question_id JOIN criteria c ON c.id=qu.criterion_id "
           "WHERE e.engagement_id=? ORDER BY e.code", (eid,))
    for x in ev:
        x["apa_intext"] = cmap.get(x["code"], {}).get("intext", "")
    ws = wb.active; ws.title = "Evidence Repository"
    sheet(ws, ["code", "comparator", "criterion", "claim", "snapshot", "author", "publisher", "title", "pub_date", "url", "locator",
               "source_category", "priority", "status", "apa_intext", "review_comment", "accessed_at"], ev)

    m = scoring.matrix(eid)
    ws = wb.create_sheet("Comparison Matrix")
    rows = [[c["name"], bm.ASSESSMENT.get(c["assessment_type"], (c["assessment_type"],))[0], "Yes" if c["is_scored"] else "No"]
            + [m["cells"][f"{c['id']}:{p['id']}"]["display"] for p in m["comparators"]] for c in m["criteria"]]
    if m["scoring"]:
        rows.append(["Overall score", "", ""] + [m["scores"][p["id"]]["overall"] for p in m["comparators"]])
        rows.append(["Coverage %", "", ""] + [m["scores"][p["id"]]["coverage"] for p in m["comparators"]])
    sheet(ws, ["Criterion", "Assessment method", "Scored"] + [p["name"] for p in m["comparators"]], rows)

    ws = wb.create_sheet("Content")
    sheet(ws, ["id", "section", "content_type", "title", "text", "evidence_ids", "status", "version"],
          q("SELECT * FROM content_items WHERE engagement_id=? ORDER BY section, id", (eid,)))
    ws = wb.create_sheet("References (APA)")
    sheet(ws, ["ref", "intext", "codes"], [{**r, "codes": ", ".join(r["codes"])} for r in references(eid, used_only=False)])
    ws = wb.create_sheet("QA Issues")
    sheet(ws, ["rule", "severity", "message", "status", "source"], q("SELECT * FROM qa_issues WHERE engagement_id=?", (eid,)))
    ws = wb.create_sheet("Search Log")
    sheet(ws, ["task_id", "query", "created_at"], q("SELECT s.* FROM search_log s JOIN research_tasks t ON t.id=s.task_id WHERE t.engagement_id=?", (eid,)))
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
